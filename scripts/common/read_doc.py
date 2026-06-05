#!/usr/bin/env python3
"""Read document content from WeChat file attachments (PDF/docx/xlsx/txt/md/csv).

Mirrors wecom's read_doc for the 'openfile' feature. Attachments under msg/file/ are plaintext,
so this works with no decryption key. Heavy parsers (pdfminer/python-docx/openpyxl) load lazily.
Usage: python read_doc.py <file> [limit]
"""
import os
import sys


def read_file(path, limit=2000):
    ext = path.lower().rsplit(".", 1)[-1] if "." in os.path.basename(path) else ""
    try:
        if ext in ("txt", "md", "csv", "log", "json", "xml", "html"):
            return open(path, encoding="utf-8", errors="replace").read()[:limit]
        if ext == "pdf":
            from pdfminer.high_level import extract_text
            return extract_text(path)[:limit]
        if ext == "docx":
            import docx
            return "\n".join(p.text for p in docx.Document(path).paragraphs)[:limit]
        if ext in ("xlsx", "xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                rows.append(f"[sheet: {ws.title}]")
                for r in ws.iter_rows(values_only=True):
                    rows.append("\t".join("" if c is None else str(c) for c in r))
                    if sum(len(x) for x in rows) > limit:
                        break
                if sum(len(x) for x in rows) > limit:
                    break
            return "\n".join(rows)[:limit]
        return f"[unsupported .{ext}; {os.path.getsize(path)} bytes -- export raw to open externally]"
    except ImportError as e:
        return f"[missing parser for .{ext}: {e}. pip install pdfminer.six python-docx openpyxl]"
    except Exception as e:
        return f"[read failed: {e}]"


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("usage: read_doc.py <file> [limit]")
    lim = int(sys.argv[2]) if len(sys.argv) > 2 else 2000
    print(read_file(sys.argv[1], lim))
